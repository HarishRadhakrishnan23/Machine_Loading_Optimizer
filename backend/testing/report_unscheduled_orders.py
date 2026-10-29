"""
report_unscheduled_orders.py - Excel list of the open WIP orders that have NOTHING schedulable left
(every open operation is QA inspection, CT = 0 / non-machine work, etc.), with their WIP details.

Read-only. Output: backend/unscheduled_orders.xlsx
Run: backend/venv/Scripts/python.exe backend/testing/report_unscheduled_orders.py
"""

import sys
import warnings
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
warnings.filterwarnings("ignore")

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db import read_wip_orders
from engine2 import service
from engine2.runner import build_chains

OUT = Path(__file__).parent.parent / "unscheduled_orders.xlsx"

REASON_TEXT = {
    "excluded_qainsp": "QA inspection (manual gate)",
    "excluded_no_routing": "No machine routing for this task (non-machine / vendor step)",
    "excluded_pn10": "Class PN10 (excluded)",
    "excluded_no_machine_for_combo": "No master row for this combination",
    "excluded_no_availability_data": "No availability data for machine",
    "excluded_unknown_fixture_device": "Fixture/locator device not in inventory",
    "excluded_no_fnl_master_ent": "No master row (fixture-required task)",
    "excluded_no_fnl_master_entry": "No master row (fixture-required task)",
}


def main() -> None:
    wip = read_wip_orders()
    wip["QC"] = wip["QUANTITY_COMPLETED"].fillna(0)
    wip["QR"] = wip["QUANTITY_REJECTED"].fillna(0)
    wip["BAL"] = wip["QUANTITY_ORDERED"] - wip["QC"] - wip["QR"]
    snap = service.get_snapshot()
    chains = build_chains(snap)
    in_plan = {c.production_order for c in chains if c.schedulable}
    open_ = wip[wip["BAL"] > 0]
    unsched = sorted(set(open_["PRODUCTION_ORDER"]) - in_plan)

    outcome = {}
    for c in chains:
        for e in c.excluded:
            outcome[(e.production_order, e.operation_no)] = e.scope_outcome.value

    last = wip.sort_values("OPERATION").groupby("PRODUCTION_ORDER").last()
    rows, detail = [], []
    for po in unsched:
        o = open_[open_["PRODUCTION_ORDER"] == po].sort_values("OPERATION")
        nxt = o.iloc[0]
        reasons = Counter()
        for _, r in o.iterrows():
            code = outcome.get((po, r["OPERATION"]))
            if code is None:
                code = "excluded_ct_zero" if (r["CYCLE_TIME"] or 0) == 0 else "unclassified"
            text = REASON_TEXT.get(code, "CT = 0 / non-machine step" if code == "excluded_ct_zero" else code)
            reasons[text] += 1
            detail.append([po, r["OPERATION"], r["TASK"], r["OPERATION_STATUS"], r["WORK_CENTER"], r["CYCLE_TIME"], r["BAL"], text])
        f = o.iloc[0]
        rows.append([
            po, f["ITEM"], f["ITEM_DESCRIPTION"], f["SIZE_INCH"], f["CLASS"], f["MOC"], f["DESIGN"], f["ORDER_STATUS"],
            f["CDD"], f["QUANTITY_ORDERED"], int(last.loc[po, "BAL"]), len(o),
            nxt["OPERATION"], nxt["TASK"], nxt["WORK_CENTER"], nxt["OPERATION_STATUS"],
            "; ".join(f"{t} x{n}" for t, n in reasons.most_common()),
        ])

    heads = ["Production Order", "Item", "Description", "Size", "Class", "MOC", "Design", "Order Status", "CDD",
             "Qty Ordered", "Pieces still to pass final op", "Open ops left", "Next open Op", "Next open Task",
             "Next open Work Center", "Next op status", "Why not scheduled"]
    df = pd.DataFrame(rows, columns=heads)
    dd = pd.DataFrame(detail, columns=["Production Order", "Operation", "Task", "Op status", "Work Center", "Cycle Time", "Balance Qty", "Why not scheduled"])
    summ = pd.DataFrame(Counter(r[-1] for r in detail).most_common(), columns=["Reason (open operation rows)", "Rows"])

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        df.to_excel(xw, "Unscheduled orders", index=False)
        dd.to_excel(xw, "Open operations", index=False)
        summ.to_excel(xw, "Summary", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F2A44")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(ws.columns, 1):
                w = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 60)
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions

    print(f"{len(df)} orders, pieces at final op: {int(df['Pieces still to pass final op'].sum())}, open op rows: {len(dd)}")
    print(summ.to_string(index=False))
    print(OUT)


if __name__ == "__main__":
    main()
