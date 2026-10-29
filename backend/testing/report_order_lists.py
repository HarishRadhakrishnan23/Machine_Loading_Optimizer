"""Read-only: two plain lists (distinct orders in MCH_WIP; distinct scheduled orders in MCH_SCHEDULE_OUTPUT) with quantities.
Output: backend/order_lists.xlsx"""
import sys, warnings
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent)); warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl.styles import Font, PatternFill
from db import get_connection, read_wip_orders

w = read_wip_orders()
w["BAL"] = w["QUANTITY_ORDERED"] - w["QUANTITY_COMPLETED"].fillna(0) - w["QUANTITY_REJECTED"].fillna(0)
last = w.sort_values("OPERATION").groupby("PRODUCTION_ORDER").last()
wip = pd.DataFrame({"Production Order": last.index, "Quantity Ordered": last["QUANTITY_ORDERED"].values,
                    "Pieces still to go (final op)": last["BAL"].clip(lower=0).values})
with get_connection() as c:
    o = pd.read_sql("SELECT DISTINCT PRODUCTION_ORDER FROM MCH_SCHEDULE_OUTPUT WHERE WORK_CENTER IS NOT NULL", c)
sched = wip[wip["Production Order"].isin(set(o["PRODUCTION_ORDER"]))].reset_index(drop=True)
with pd.ExcelWriter(Path(__file__).parent.parent / "order_lists.xlsx", engine="openpyxl") as xw:
    wip.to_excel(xw, "In WIP", index=False); sched.to_excel(xw, "Scheduled", index=False)
    for ws in xw.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="2B0C53")
        for col, wd in zip("ABC", (20, 18, 30)): ws.column_dimensions[col].width = wd
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
print(len(wip), wip.iloc[:, 1].sum(), wip.iloc[:, 2].sum()); print(len(sched), sched.iloc[:, 1].sum(), sched.iloc[:, 2].sum())
