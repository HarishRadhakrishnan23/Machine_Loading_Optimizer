"""
report_missing_fnl_master_rows.py — Excel list of the (SIZE, CLASS, MOC, DESIGN, TASK) combinations that have
open WIP operations but NO row in MCH_MACHINE_FNL_MASTER (the engine classifies them
EXCLUDED_NO_MACHINE_FOR_COMBO and leaves them out of the plan), laid out in the master-data format so
Engineering can fill in Priority / Work_Center / fixture / locator / times.

Read-only: never writes to Oracle. Output: backend/missing_fnl_master_rows.xlsx

Run: backend/venv/Scripts/python.exe backend/testing/report_missing_fnl_master_rows.py
"""

import sys
import warnings
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
warnings.filterwarnings("ignore")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine2 import service
from engine2.runner import build_chains

OUT = Path(__file__).parent.parent / "missing_fnl_master_rows.xlsx"

# Exactly the layout Engineering asked for (first six are pre-filled; the rest are theirs to complete).
HEADERS = [
    "Size", "Class", "MOC", "Design", "Combination", "Task", "Priority", "Work_Center", "Fixture_ID", "Locator_ID",
    "Fixture_Change_Time", "Locator Change", "Clamping Set up", "Tooling Change", "First Piece Prooving",
    "Locator_Change_Time", "Loading_Unloading_Time",
]
PREFILLED = 6


def num(v: str):
    """'14' -> 14, '150' -> 150, '8.5' -> 8.5, anything else unchanged (e.g. 'PN10')."""
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


def main() -> None:
    snap = service.get_snapshot()
    chains = build_chains(snap)
    in_plan = {c.production_order for c in chains if c.schedulable}
    typed = {(r.production_order, r.operation_no): r for r in snap.raw_rows}  # typed SIZE/CLASS/MOC/DESIGN — never parsed from a key

    stats: dict[tuple, dict] = defaultdict(lambda: {"orders": set(), "ops": 0, "pieces": 0, "in_plan": set()})
    for c in chains:
        for e in c.excluded:
            if e.scope_outcome.value != "excluded_no_machine_for_combo":
                continue
            r = typed[(e.production_order, e.operation_no)]
            key = (r.size_inch, r.class_val, r.moc, r.design, e.task)
            s = stats[key]
            s["orders"].add(e.production_order)
            s["ops"] += 1
            s["pieces"] += e.balance_qty
            if e.production_order in c.production_order and c.production_order in in_plan:
                s["in_plan"].add(c.production_order)

    def sort_key(k):
        size, cls, moc, design, task = k
        n = num(size)
        return (task, n if isinstance(n, (int, float)) else 1e9, str(cls), str(moc), str(design))

    keys = sorted(stats, key=sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Missing master rows"
    head_fill = PatternFill("solid", fgColor="1F2A44")
    known_fill = PatternFill("solid", fgColor="E8EDF5")
    todo_fill = PatternFill("solid", fgColor="FFF6D6")
    ws.append(HEADERS)
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill if i <= PREFILLED else PatternFill("solid", fgColor="8A5A00")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for size, cls, moc, design, task in keys:
        combination = f"{size}-{cls}-{moc}-{design}"  # same text convention as MCH_MACHINE_FNL_MASTER.COMBINATION
        ws.append([num(size), num(cls), moc, design, combination, task] + [None] * (len(HEADERS) - PREFILLED))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.fill = known_fill if c.column <= PREFILLED else todo_fill
    widths = [8, 8, 18, 9, 28, 8, 9, 14, 24, 24, 18, 16, 16, 15, 20, 20, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # ── impact sheet: how much each missing row matters, so Engineering can prioritise ──
    wi = wb.create_sheet("Impact (prioritise)")
    wi.append(["Size", "Class", "MOC", "Design", "Task", "Orders affected", "Of those, orders in today's plan", "Open operation rows", "Pieces waiting at this operation"])
    for i in range(1, 10):
        c = wi.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k in sorted(keys, key=lambda k: (-len(stats[k]["orders"]), -stats[k]["pieces"])):
        s = stats[k]
        wi.append([num(k[0]), num(k[1]), k[2], k[3], k[4], len(s["orders"]), len(s["in_plan"] & s["orders"]), s["ops"], s["pieces"]])
    for i, w in enumerate([8, 8, 18, 9, 8, 15, 22, 18, 24], 1):
        wi.column_dimensions[get_column_letter(i)].width = w
    wi.row_dimensions[1].height = 44
    wi.freeze_panes = "A2"
    wi.auto_filter.ref = f"A1:I{wi.max_row}"

    # ── notes ──
    wn = wb.create_sheet("Notes")
    notes = [
        ("What this is", "Combinations (Size / Class / MOC / Design) + Task that have open work orders in WIP but no row in MCH_MACHINE_FNL_MASTER."),
        ("Effect today", "The scheduler cannot place those operations on any machine, so it skips them. The affected orders' planned completion dates ignore that work and are optimistic."),
        ("Pre-filled", "Size, Class, MOC, Design, Combination, Task (grey-blue columns). Combination uses the master's own text style: Size-Class-MOC-Design."),
        ("For Engineering to fill (yellow)", "Priority (1 = most preferred machine), Work_Center, Fixture_ID, Locator_ID and the times, in MINUTES."),
        ("One row per machine", "A combination that can run on several machines needs one row per machine (copy the pre-filled columns); Fixture_ID and Locator_ID stay the same across machines of one combination+task, times may differ per machine."),
        ("Priority", "Unique per combination+task: 1 = preferred, 2 = next, ..."),
        ("Fixture / locator", "Use a device name from MCH_FIXTURE_LOCATOR. 'NA' = no physical device needed. A '+'-joined locator (A+B) means both devices are needed together."),
        ("Columns the system does not read yet", "Locator Change, Clamping Set up, Tooling Change and First Piece Prooving are NOT columns of MCH_MACHINE_FNL_MASTER today. The scheduler uses only Fixture_Change_Time, Locator_Change_Time and Loading_Unloading_Time. Anything entered in the other four is kept for reference until the table and the scheduler are extended."),
        ("Impact sheet", "Orders affected / pieces waiting per missing row, biggest first - fill the top rows first for the largest effect."),
    ]
    wn.column_dimensions["A"].width = 34
    wn.column_dimensions["B"].width = 120
    for a, b in notes:
        wn.append([a, b])
    for row in wn.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    print(f"{len(keys)} missing combination+task rows -> {OUT}")
    print(f"orders affected: {len({o for s in stats.values() for o in s['orders']})}, "
          f"in today's plan: {len({o for s in stats.values() for o in s['in_plan']})}, "
          f"open operation rows: {sum(s['ops'] for s in stats.values())}")


if __name__ == "__main__":
    main()
