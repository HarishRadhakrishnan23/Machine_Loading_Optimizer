"""
report_wip_vs_schedule_orders.py - Excel with two sheets: distinct orders (and quantities) in MCH_WIP vs in
MCH_SCHEDULE_OUTPUT, plus a comparison of the two.

Read-only. Output: backend/wip_vs_schedule_orders.xlsx
Run: backend/venv/Scripts/python.exe backend/testing/report_wip_vs_schedule_orders.py
"""

import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
warnings.filterwarnings("ignore")

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db import get_connection, read_wip_orders

OUT = Path(__file__).parent.parent / "wip_vs_schedule_orders.xlsx"


def main() -> None:
    wip = read_wip_orders()
    wip["BAL"] = wip["QUANTITY_ORDERED"] - wip["QUANTITY_COMPLETED"].fillna(0) - wip["QUANTITY_REJECTED"].fillna(0)
    wip = wip.sort_values(["PRODUCTION_ORDER", "OPERATION"])
    g = wip.groupby("PRODUCTION_ORDER")
    wip_orders = pd.DataFrame({
        "Order Status": g["ORDER_STATUS"].first(),
        "Item": g["ITEM"].first(),
        "CDD": g["CDD"].first(),
        "Qty Ordered": g["QUANTITY_ORDERED"].max(),
        "Operations (rows)": g.size(),
        "Open operations (balance > 0)": g["BAL"].apply(lambda s: int((s > 0).sum())),
        "Balance qty - final operation": g["BAL"].last(),
        "Balance qty - sum over all operations": g["BAL"].apply(lambda s: s.clip(lower=0).sum()),
    }).reset_index().rename(columns={"PRODUCTION_ORDER": "Production Order"})

    with get_connection() as conn:
        out = pd.read_sql(
            "SELECT PRODUCTION_ORDER, OPERATION_NO, LINE_NO, BALANCE_QTY, WORK_CENTER FROM MCH_SCHEDULE_OUTPUT", conn
        )
    out.columns = [c.upper() for c in out.columns]
    out = out.sort_values(["PRODUCTION_ORDER", "OPERATION_NO", "LINE_NO"])
    # one operation can be split over several LINE_NO slices: its quantity is the sum of the slices
    per_op = out.groupby(["PRODUCTION_ORDER", "OPERATION_NO"], as_index=False).agg(
        qty=("BALANCE_QTY", "sum"), scheduled=("WORK_CENTER", lambda s: s.notna().any())
    )
    h = per_op.sort_values("OPERATION_NO").groupby("PRODUCTION_ORDER")
    sched_orders = pd.DataFrame({
        "Operation rows (incl. unscheduled)": h.size(),
        "Operations placed on a machine": h["scheduled"].sum().astype(int),
        "Qty - final operation": h["qty"].last(),
        "Qty - sum over all operations": h["qty"].sum(),
    })
    sched_orders["Output rows (incl. slices)"] = out.groupby("PRODUCTION_ORDER").size()
    sched_orders = sched_orders.reset_index().rename(columns={"PRODUCTION_ORDER": "Production Order"})

    w_set, s_set = set(wip_orders["Production Order"]), set(sched_orders["Production Order"])
    summary = pd.DataFrame([
        ["Distinct orders in MCH_WIP", len(w_set), int(wip_orders["Qty Ordered"].sum()),
         int(wip_orders["Balance qty - final operation"].clip(lower=0).sum())],
        ["Distinct orders in MCH_SCHEDULE_OUTPUT", len(s_set), None, int(sched_orders["Qty - final operation"].sum())],
        ["In WIP but not in MCH_SCHEDULE_OUTPUT", len(w_set - s_set), None, None],
        ["In MCH_SCHEDULE_OUTPUT but not in WIP", len(s_set - w_set), None, None],
    ], columns=["Measure", "Distinct orders", "Total Qty Ordered", "Total balance qty (final operation)"])
    notes = pd.DataFrame({"Notes": [
        "Quantity is per order: 'Qty Ordered' is the order size; 'final operation' balance is the pieces still to pass the last op.",
        "'Sum over all operations' counts the same pieces once per operation, so it is larger - shown for reference only.",
        "MCH_SCHEDULE_OUTPUT quantities add up the LINE_NO slices of an operation (a split operation is not counted twice).",
        "Balance qty = QUANTITY_ORDERED - QUANTITY_COMPLETED - QUANTITY_REJECTED.",
    ]})

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        summary.to_excel(xw, "Summary", index=False)
        notes.to_excel(xw, "Summary", index=False, startrow=len(summary) + 3)
        wip_orders.to_excel(xw, "WIP orders", index=False)
        sched_orders.to_excel(xw, "Schedule output orders", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2B0C53")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(ws.columns, 1):
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 12), 60)
            ws.freeze_panes = "A2"
            if ws.title != "Summary":
                ws.auto_filter.ref = ws.dimensions
    print(summary.to_string(index=False))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
